from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
TEMPLATE_PATH = BASE_DIR / "template_2026.xlsx"


def make_template() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    wb.create_sheet("数据释意")
    # create minimal grid labels
    for r in range(1, 45):
        ws.cell(row=r, column=1, value=f"row_{r}")
    wb.save(TEMPLATE_PATH)


def make_inputs() -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)

    pd.DataFrame(
        [
            {
                "order_id": "SO001",
                "order_date": "2026-01-05",
                "settlement_month": "2026-01",
                "channel": "淘宝",
                "product_name": "商品A",
                "sku": "SKU001",
                "qty": 2,
                "gross_sales_amount": "¥1,200.00",
                "refund_amount": "100",
                "net_sales_amount": "1100",
            },
            {
                "order_id": "SO002",
                "order_date": "2026-02-10",
                "settlement_month": "2月",
                "channel": "天猫",
                "product_name": "商品B",
                "sku": "SKU002",
                "qty": 1,
                "gross_sales_amount": "800",
                "refund_amount": "0",
                "net_sales_amount": "800",
            },
        ]
    ).to_csv(INPUT_DIR / "sales_orders.csv", index=False)

    pd.DataFrame(
        [
            {
                "txn_id": "IN001",
                "received_date": "2026-01-20",
                "settlement_month": "2026/01",
                "income_type": "渠道分销",
                "counterparty": "CP1",
                "amount": "300",
            },
            {
                "txn_id": "IN002",
                "received_date": "2026-01-25",
                "settlement_month": "1",
                "income_type": "IP返采",
                "counterparty": "CP2",
                "amount": "500",
            },
        ]
    ).to_csv(INPUT_DIR / "income_statement.csv", index=False)

    pd.DataFrame(
        [
            {"settlement_month": "1", "ip_name": "战双", "royalty_amount": "100"},
            {"settlement_month": "01", "ip_name": "ES", "royalty_amount": "120"},
        ]
    ).to_csv(INPUT_DIR / "royalty_details.csv", index=False)

    pd.DataFrame(
        [
            {"cost_id": "C001", "settlement_month": "1", "cost_type": "product_cost", "amount": "600"},
            {"cost_id": "C002", "settlement_month": "1", "cost_type": "warehouse_cost", "amount": "80"},
            {"cost_id": "C003", "settlement_month": "1", "cost_type": "shipping_cost", "amount": "70"},
        ]
    ).to_csv(INPUT_DIR / "cost_details.csv", index=False)

    pd.DataFrame(
        [
            {"expense_id": "E001", "settlement_month": "1", "expense_type": "customer_service_base_fee", "amount": "60"},
            {"expense_id": "E002", "settlement_month": "1", "expense_type": "customer_service_bonus", "amount": "20"},
            {"expense_id": "E003", "settlement_month": "1", "expense_type": "store_operation_base_fee", "amount": "30"},
            {"expense_id": "E004", "settlement_month": "1", "expense_type": "store_misc_software_fee", "amount": "10"},
            {"expense_id": "E005", "settlement_month": "1", "expense_type": "store_operation_bonus", "amount": "15"},
        ]
    ).to_csv(INPUT_DIR / "operation_expenses.csv", index=False)

    pd.DataFrame(
        [
            {
                "fee_id": "F001",
                "settlement_month": "1",
                "fee_type": "ad_spend",
                "platform": "淘宝",
                "amount": "200",
            },
            {
                "fee_id": "F002",
                "settlement_month": "1",
                "fee_type": "platform_commission_fee",
                "platform": "淘宝",
                "amount": "90",
            },
        ]
    ).to_csv(INPUT_DIR / "platform_fees.csv", index=False)


if __name__ == "__main__":
    make_template()
    make_inputs()
    print(f"Generated template: {TEMPLATE_PATH}")
    print(f"Generated input dir: {INPUT_DIR}")
