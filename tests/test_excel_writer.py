from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.excel_writer import write_back_to_excel
from src.models import ValidationReport


def _create_template(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    wb.create_sheet("数据释意")
    wb.save(path)


def test_excel_write_and_formulas(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    _create_template(template)

    aggregated = {m: {} for m in range(1, 13)}
    aggregated[1] = {"ecom_gross_sales": 1000.0, "ecom_net_sales_after_refund": 900.0}

    write_back_to_excel(str(template), str(output), 2026, aggregated, ValidationReport())

    wb = load_workbook(output)
    ws = wb["2026"]

    assert ws["D2"].value == 1000.0
    assert ws["D5"].value == 900.0
    assert ws["D39"].value == "=D21+D23+D25+D26+D33+D37"
    assert ws["D40"].value == "=D6-D39"
    assert ws["M40"].value == "=M6-M39"
    assert ws["N39"].value == "=N21+N23+N25+N26+N33+N37"
    assert ws["O43"].value == "=IF(O6=0,0,O42/O6)"

    assert "validation_report" in wb.sheetnames
