from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from config.mappings import MONTH_TO_COLUMN, SHEET_ROW_MAPPING
from config.settings import SETTINGS
from src.models import ValidationReport


def _ratio_formula(numerator_ref: str, denominator_ref: str) -> str:
    if SETTINGS.ratio_default_zero:
        return f"=IF({denominator_ref}=0,0,{numerator_ref}/{denominator_ref})"
    return f"=IF({denominator_ref}=0,\"\",{numerator_ref}/{denominator_ref})"


def write_back_to_excel(
    template_path: str,
    output_path: str,
    year: int,
    aggregated: dict[int, dict[str, float]],
    report: ValidationReport,
) -> None:
    wb = load_workbook(template_path)
    sheet_name = str(year)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Main sheet not found: {sheet_name}")
    ws = wb[sheet_name]

    for month in range(1, 13):
        col = MONTH_TO_COLUMN[month]
        month_data = aggregated.get(month, {})
        for metric, row in SHEET_ROW_MAPPING.items():
            ws[f"{col}{row}"] = float(month_data.get(metric, 0.0))

        ws[f"{col}6"] = f"={col}5+{col}3+{col}4"
        ws[f"{col}7"] = f"={col}6/(1+{SETTINGS.tax_rate_income})"
        ws[f"{col}21"] = f"=SUM({col}8:{col}20)"
        ws[f"{col}22"] = _ratio_formula(f"{col}21", f"{col}6")
        ws[f"{col}24"] = _ratio_formula(f"{col}23", f"{col}6")
        ws[f"{col}27"] = _ratio_formula(f"{col}25+{col}26", f"{col}6")
        ws[f"{col}33"] = f"=SUM({col}28:{col}32)"
        ws[f"{col}34"] = _ratio_formula(f"{col}33", f"{col}6")
        ws[f"{col}37"] = f"={col}35+{col}36"
        ws[f"{col}38"] = _ratio_formula(f"{col}37", f"{col}6")
        ws[f"{col}39"] = f"={col}21+{col}23+{col}25+{col}26+{col}33+{col}37"
        ws[f"{col}40"] = f"={col}6-{col}39"
        ws[f"{col}41"] = _ratio_formula(f"{col}40", f"{col}6")
        ws[f"{col}42"] = f"={col}6-{col}23"
        ws[f"{col}43"] = _ratio_formula(f"{col}42", f"{col}6")

    if SETTINGS.validation_sheet_name in wb.sheetnames:
        del wb[SETTINGS.validation_sheet_name]
    report_ws = wb.create_sheet(SETTINGS.validation_sheet_name)
    report_df = report.to_dataframe()
    headers = list(report_df.columns)
    for col_idx, head in enumerate(headers, start=1):
        report_ws.cell(row=1, column=col_idx, value=head)
    for row_idx, row in enumerate(report_df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            report_ws.cell(row=row_idx, column=col_idx, value=value)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
